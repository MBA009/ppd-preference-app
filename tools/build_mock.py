#!/usr/bin/env python3
"""Build the static JSON fixtures in public/mock/ from the research workbook.

    python3 tools/build_mock.py [path/to/Dataflow_updated_v3.xlsx]

The app never reads the workbook directly. It reads the JSON written here through
src/api/mockClient.ts, so a real backend only has to serve the same shapes
(documented in src/api/types.ts).

Each flow's taxonomy uses only the core ppd: terms of draft-dsmullen-ppd-taxonomy-05 (the workbook's
"Core rule groups"), with constraints as the protocol's Constraints Object. The workbook's finer-grained
research terms only feed the plain-language labels. The fixtures do NOT hold verdicts: a verdict compares
a flow with the household's rules; app-api computes it on read with shared/engine.ts.
"""
import json, os, re, sys, collections
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "public", "mock")
XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "Dataflow_updated_v3.xlsx")

# ---------------------------------------------------------------- vocabulary
def words(iri):
    s = iri.split(":")[-1]
    s = re.sub(r"(?<!^)(?=[A-Z])", " ", s).lower()
    return s[:1].upper() + s[1:]

DATA_TYPE = {
    "aiPromptData": "Questions you ask the AI assistant",
    "alertEventData": "Alert and event details",
    "audienceSegment": "The audience group you are placed in",
    "audioData": "Sound from the microphone",
    "audioFingerprint": "A pattern of the sound that is playing",
    "connectedDeviceInfo": "Devices connected to the TV",
    "cryingDetectionEvent": "Crying sound detected",
    "deviceControlCommand": "Commands to your devices",
    "deviceCredentialData": "A password for the device",
    "diagnosticData": "Device logs and status",
    "environmentData": "Temperature, humidity or light readings",
    "geofenceState": "Whether you are home or away",
    "hdmiDeviceInfo": "Devices connected by HDMI",
    "householdDeviceState": "The state of your home devices",
    "householdMemberAccessGrant": "Who you give access to",
    "hvacSystemProfileData": "Your heating and cooling setup",
    "lockControlCommand": "Lock and unlock commands",
    "lockState": "Whether the door is locked",
    "nearbyWifiData": "Nearby Wi-Fi networks",
    "networkCredentialData": "Your Wi-Fi password",
    "networkTelemetry": "Wi-Fi connection details",
    "pairingCredentialData": "A pairing code",
    "presenceData": "Whether someone is there",
    "thermostatControlCommand": "Thermostat commands",
    "visualData": "Video or photos",
    "voiceRecognitionResult": "What the TV heard you say",
    "voiceTechnicalLog": "Logs from the voice feature",
    "wifiNetworkName": "Your Wi-Fi network name",
    "deviceAndNetworkData": "Device model and network details",
    "identifierData": "Names, accounts or device IDs",
    "locationData": "Location",
    "usageData": "How you use the device",
}
# (label, short label used under a row title)
PURPOSE = {
    "accessControl": ("Controlling who can get in", "access"),
    "accessSharing": ("Sharing access with other people", "sharing access"),
    "accountSetup": ("Setting up your account", "account setup"),
    "acrLocationSupport": ("Supporting viewing recognition with location", "location"),
    "activityHistory": ("Showing activity history", "activity history"),
    "adaptiveClimate": ("Adjusting the climate automatically", "climate"),
    "aiAssistant": ("AI assistant", "AI assistant"),
    "audienceMeasurement": ("Measuring who watches", "audience measurement"),
    "autoUnlock": ("Unlocking when you arrive", "auto unlock"),
    "cameraAdaptation": ("Adjusting the camera", "camera"),
    "cloudStorage": ("Cloud storage", "cloud storage"),
    "contentRecognition": ("Recognizing what is on screen", "recognition"),
    "customerSupport": ("Customer support", "support"),
    "deviceDiscovery": ("Finding connected devices", "discovery"),
    "deviceSetup": ("Setting up the device", "setup"),
    "displayAdaptation": ("Adjusting the picture", "picture"),
    "ecosystemIntegration": ("Working with your smart-home app", "integration"),
    "emailAlert": ("Email alerts", "email alerts"),
    "hvacControl": ("Heating and cooling", "climate"),
    "localConnectivity": ("Connecting to nearby phones", "nearby connection"),
    "locationBasedMarketing": ("Location-based marketing", "marketing"),
    "occupancyControl": ("Detecting if someone is home", "presence"),
    "pushAlert": ("Push alerts", "push alerts"),
    "remoteAccessControl": ("Controlling the lock away from home", "remote access"),
    "remoteControl": ("Controlling the device away from home", "remote control"),
    "remoteViewing": ("Watching remotely", "remote viewing"),
    "researchAndDevelopment": ("Product research", "research"),
    "smartDisplayViewing": ("Showing video on a smart display", "smart display"),
    "smartHomeControl": ("Controlling other devices", "control"),
    "twoWayAudio": ("Talking through the camera", "two-way audio"),
    "usageEfficiency": ("Product use", "product use"),
    "visitorViewing": ("Showing who is at the door", "visitors"),
    "voiceControl": ("Voice control", "voice control"),
    "advertisingAndMarketing": ("Advertising and marketing", "advertising"),
    "diagnosticsAndMaintenance": ("Diagnostics and maintenance", "diagnostics"),
    "security": ("Security", "security"),
}
ACTION = {"collection": "Collected", "transfer": "Sent to another party", "use": "Used", "inference": "Worked out from other data"}
SOURCE = {
    "userProvided": "You", "participantGenerated": "The device", "participantObserved": "Observed by the device",
    "derivedFromPriorData": "Worked out from earlier data", "vendorProvided": "The vendor", "householdProvided": "Your household",
    "linkedClimateSensor": "A linked climate sensor", "smartHomeEcosystemProvided": "Your smart-home app",
    "mmWaveSensor": "The built-in presence sensor", "pairedDoorbell": "A paired doorbell", "lockMechanism": "The lock",
    "mobileLocationService": "Your phone's location", "ambientLightSensor": "The light sensor", "cameraSensor": "The camera",
    "microphone": "The microphone", "daylightSensor": "The daylight sensor",
}
# (label, short)
HANDLER = {
    "householdContext": ("Your household", "your home"), "aqaraService": ("Aqara", "Aqara"), "levelService": ("Level", "Level"),
    "lgService": ("LG", "LG"), "reolinkService": ("Reolink", "Reolink"), "reolinkCloud": ("Reolink Cloud", "Reolink Cloud"),
    "advertisingService": ("An advertising service", "an ad service"), "advertiser": ("Advertisers", "advertisers"),
    "thirdPartyServiceProvider": ("A third-party service", "partners"),
    "smartHomeEcosystem": ("Your smart-home app (Alexa, Google…)", "your smart-home app"),
    "voiceRecognitionService": ("A voice-recognition service", "a voice service"),
    "acrService": ("LG's viewing-recognition service", "LG's recognition service"),
}
CONTROL_TYPE = {
    "user_setting": "Change in settings", "setup_step": "Provided during setup", "user_initiated": "Happens only when you do it",
    "external_request": "Only by contacting the vendor", "none_documented": "No control found", "unknown": "Not checked, or sources silent",
}
CORE_DT = {"contentData": "Content (video, audio, photos)", "deviceAndNetworkData": "Device and network details", "identifierData": "Identifiers and accounts",
           "locationData": "Location", "profileData": "Profile information", "sensorData": "Sensor readings", "usageData": "Usage data"}
CORE_ACTION_HANDLER = {
    ("transfer", "vendorContext"): "sent to a vendor", ("transfer", "thirdPartyContext"): "sent to a third party", ("transfer", "householdContext"): "sent within your home",
    ("collection", "vendorContext"): "collected by a vendor", ("collection", "thirdPartyContext"): "collected by a third party", ("collection", "householdContext"): "collected at home",
    ("use", "vendorContext"): "used by a vendor", ("use", "thirdPartyContext"): "used by a third party", ("use", "householdContext"): "used at home",
    ("inference", "vendorContext"): "worked out by a vendor", ("inference", "thirdPartyContext"): "worked out by a third party", ("inference", "householdContext"): "worked out at home",
}

def local(iri):
    return iri.split(":")[-1] if iri else None

# ---------------------------------------------------------------- titles
TITLES = [  # (device id, lower-case substring of the workbook "Dataflow" text, title)
    ("aqara-w200", "analytics/advertising", "Usage shared with ad partners"),
    ("aqara-w200", "precise mobile geolocation", "Phone location for marketing"),
    ("aqara-w200", "device diagnostics / logs", "Device logs sent to Aqara"),
    ("aqara-w200", "wi-fi / network technical telemetry", "Wi-Fi details sent to Aqara"),
    ("aqara-w200", "device usage history", "Usage history sent to Aqara"),
    ("reolink-e1-zoom", "cloud video upload", "Cloud video upload"),
    ("reolink-e1-zoom", "sent for push notifications", "Push alerts through Reolink"),
    ("reolink-e1-zoom", "live / recorded video delivered", "Remote viewing through Reolink"),
    ("reolink-e1-zoom", "shared to alexa", "Smart-display sharing"),
    ("reolink-e1-zoom", "app usage analytics", "App analytics"),
    ("lg-oled65c6pua-c6", "shared with advertiser", "Viewing data shared with advertisers"),
    ("lg-oled65c6pua-c6", "voice audio transferred", "Voice recordings sent for recognition"),
    ("lg-oled65c6pua-c6", "smart-tv device / software", "TV diagnostics sent to LG"),
]

def title_for(dev, text):
    low = text.lower()
    for d, sub, t in TITLES:
        if d == dev and sub in low:
            return t
    t = re.sub(r"^setup:\s*", "", text, flags=re.I)
    t = re.split(r"\s\(", t)[0].strip()
    if len(t) > 78:
        t = t[:78].rsplit(" ", 1)[0] + "…"
    return t[:1].upper() + t[1:]

# ---------------------------------------------------------------- devices (app-level facts the workbook does not hold)
DEVICE_META = {
    "aqara-w200":        {"short": "Aqara Thermostat Hub W200", "icon": "thermostat", "matter": True,  "source": "declared", "state": "waiting_for_renewal"},
    "level-lock":        {"short": "Level Lock",                "icon": "lock",       "matter": True,  "source": "advisory", "state": "advisory"},
    "lg-oled65c6pua-c6": {"short": "LG OLED65C6PUA C6",         "icon": "tv",         "matter": False, "source": "advisory", "state": "advisory"},
    "reolink-e1-zoom":   {"short": "Reolink E1 Zoom",           "icon": "camera",     "matter": False, "source": "declared", "state": "rules_confirmed"},
}
SEED_HOME = ["reolink-e1-zoom", "aqara-w200", "level-lock", "lg-oled65c6pua-c6"]

def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    rows = list(wb["Flows"].iter_rows(values_only=True))
    head = rows[0]
    flows = [dict(zip(head, r)) for r in rows[1:] if r[0]]
    grows = list(wb["Core rule groups"].iter_rows(values_only=True))
    gh = grows[0]
    groups = {}
    for r in grows[1:]:
        if r[0]:
            g = dict(zip(gh, r))
            groups[g["Core rule group"]] = {k: g[k] for k in ("data_type", "purpose", "action", "source", "handling_context")}
    drows = list(wb["Devices"].iter_rows(values_only=True))
    dh = drows[0]
    devices = [dict(zip(dh, r)) for r in drows[1:] if r[0]]

    os.makedirs(os.path.join(OUT, "devices"), exist_ok=True)
    library, codes = [], {}
    by_dev = collections.defaultdict(list)
    for f in flows:
        by_dev[f["Device ID"]].append(f)

    for d in devices:
        did = d["Device ID"]
        meta = DEVICE_META[did]
        out_flows = []
        for i, f in enumerate(by_dev[did]):
            core = groups[f["Core rule group"]]
            dt, pur, act, src, hc = (local(f[k]) for k in ("data_type", "purpose", "action", "source", "handling_context"))
            cons = {}
            for c in (f["Constraints"] or "").split(";"):
                if "=" in c:
                    k, v = c.strip().split("=")
                    cons[k] = local(v)
            kept = {"indefinite": "Until deleted", "ephemeral": "Not kept"}.get(cons.get("retention"), "Not stated")
            limit = {"onDeviceOnly": "Stays on the device", "inHomeOnly": "Stays in your home"}.get(cons.get("processing_boundary"))
            handler_label, handler_short = HANDLER[hc]
            what_happens = ACTION[act]
            if act == "transfer" and hc == "householdContext":
                what_happens = "Sent within your home"
            verb = {"transfer": "Goes to", "collection": "Collected by", "use": "Used by", "inference": "Worked out by"}[act]
            who = "Stays in your home" if hc == "householdContext" else "%s %s" % (verb, handler_short)
            summary = who + " · " + PURPOSE[pur][1]
            gl = "%s %s" % (CORE_DT.get(local(core["data_type"]), words(core["data_type"])), CORE_ACTION_HANDLER.get((local(core["action"]), local(core["handling_context"])), ""))
            out_flows.append({
                "id": f["Flow ID"],
                "title": title_for(did, f["Dataflow"]),
                "workbook_text": f["Dataflow"],
                "summary": summary,
                "summary_parts": {"who": who, "purpose": PURPOSE[pur][1]},
                "phase": f["Phase"],
                "taxonomy": dict(core, constraints={k: "ppd:" + v for k, v in cons.items()}),
                "display": {"what": DATA_TYPE.get(dt, words(dt)), "why": PURPOSE[pur][0], "what_happens": what_happens,
                            "comes_from": SOURCE.get(src, words(src)), "handled_by": handler_label, "kept": kept, "limit": limit},
                "core_group": {"id": f["Core rule group"], "label": gl.strip()},
                "control": {"type": f["Control type"], "type_label": CONTROL_TYPE[f["Control type"]],
                            "necessity": f["Necessity"], "note": f["Control note (draft)"]},
                "evidence": f["Evidence"],
                "core": core,
            })
        decl_id = d["Declaration device_id (example)"]
        declaration = None
        if meta["source"] == "declared":
            declaration = {"state": meta["state"], "declaration_id": decl_id,
                           "received_at": "2026-10-14T09:12:00Z", "renew_by": "2026-10-21T09:12:00Z",
                           "statements": d["Declaration statements"]}
        dev = {"id": did, "name": re.sub(r"\s*\(.*\)$", "", d["Device"]), "short_name": meta["short"], "category": d["Category ID"],
               "category_label": d["Category"], "icon": meta["icon"], "matter": meta["matter"], "source": meta["source"],
               "declaration": declaration, "declaration_id_example": decl_id, "flow_count": len(out_flows), "flows": out_flows}
        json.dump(dev, open(os.path.join(OUT, "devices", did + ".json"), "w"), indent=1, ensure_ascii=False)
        library.append({k: v for k, v in dev.items() if k != "flows"})
        codes["https://ppd.example/d/" + decl_id] = did

    json.dump({"devices": library}, open(os.path.join(OUT, "library.json"), "w"), indent=1, ensure_ascii=False)
    json.dump({"codes": codes,
               "note": "Placeholder privacy-code payloads. A real code would point at the device's declaration endpoint."},
              open(os.path.join(OUT, "codes.json"), "w"), indent=1)
    json.dump({"devices": [{"device_id": d, "added_at": "2026-10-14T09:15:00Z"} for d in SEED_HOME]},
              open(os.path.join(OUT, "household.json"), "w"), indent=1)

    print("wrote", len(library), "devices,", sum(d["flow_count"] for d in library), "flows")
    for l in library:
        print("  %-26s %d flows" % (l["short_name"], l["flow_count"]))

main()
